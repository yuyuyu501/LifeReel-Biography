from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from lifereel_api.modules.interview.profile_service import anonymized, value_text

STATES = {
    "empty": "待补充",
    "filled": "已有资料",
    "unknown": "暂不清楚",
    "deferred": "暂时跳过",
    "declined": "不愿回答",
    "not_applicable": "不适用",
}


def export_rows(profile, include_private=False, sections=None):
    selected = set(sections or [s["key"] for s in profile["sections"]])
    rows = []
    for field in profile["fields"]:
        if field["section"] not in selected:
            continue
        matches = [e for e in profile["entries"] if e["field_key"] == field["key"]]
        for entry in matches or [None]:
            if entry and entry["use_scope"] == "internal" and not include_private:
                continue
            rows.append((field, anonymized(entry) if entry and not include_private else entry))
    return rows


def markdown(profile, include_private=False, sections=None):
    rows = export_rows(profile, include_private, sections)
    name = next(
        (
            anonymized(e)["value"]
            for e in profile["entries"]
            if e["field_key"] == "identity.preferred_name" and e["use_scope"] != "internal"
        ),
        "人生资料",
    )
    lines = [
        f"# {name}的人生资料",
        "",
        f"资料版本：{profile['version_number']}",
        f"模板版本：{profile['template_version']}",
        f"导出日期：{datetime.now(UTC).date().isoformat()}",
        "此文件是资料快照，修改后不会自动回写网站。",
        "",
    ]
    for section in profile["sections"]:
        selected = [(f, e) for f, e in rows if f["section"] == section["key"]]
        if not selected:
            continue
        lines.append("## " + section["title"])
        for field, entry in selected:
            lines.extend(
                [
                    "",
                    "### " + field["label"],
                    value_text(entry["value"]) if entry else "待补充",
                    "状态：" + STATES.get(entry["state"], entry["state"])
                    if entry
                    else "状态：待补充",
                ]
            )
            if entry:
                source = entry["source"]
                lines.extend(
                    [
                        f"使用范围：{entry['use_scope']}",
                        f"来源：{source.get('type', '')} / {source.get('id', '')}",
                    ]
                )
    return "\n".join(lines)


def xlsx(profile, include_private=False, sections=None):
    rows = export_rows(profile, include_private, sections)
    book = Workbook()
    info = book.active
    info.title = "说明与概况"
    info.append(["项目", "内容"])
    info.append(["资料版本", str(profile["version_number"])])
    info.append(["模板版本", profile["template_version"]])
    info.append(["导出日期", datetime.now(UTC).date().isoformat()])
    info.append(["范围", "包含内部资料" if include_private else "可用于作品的资料"])
    info.append(["说明", "此文件是导出快照，修改后不会自动回写网站。未知信息不代表没有发生。"])
    sheet = book.create_sheet("资料总表")
    sheet.append(
        ["类别", "项目", "字段键", "记录ID", "当前内容", "处理状态", "确定性", "使用范围", "来源"]
    )
    events = book.create_sheet("经历明细")
    events.append(
        [
            "记录ID",
            "类别",
            "标题",
            "时间原话",
            "时间精度",
            "地点",
            "人物",
            "经过",
            "行动",
            "感受",
            "结果或影响",
        ]
    )
    people = book.create_sheet("人物与关联")
    people.append(["记录ID", "项目", "关联内容"])
    sources = book.create_sheet("来源与素材")
    sources.append(["记录ID", "来源类型", "来源ID", "来源版本", "证据摘要", "使用范围"])
    section_names = {s["key"]: s["title"] for s in profile["sections"]}
    for field, entry in rows:
        if not entry:
            sheet.append(
                [
                    section_names[field["section"]],
                    field["label"],
                    field["key"],
                    "",
                    "",
                    "待补充",
                    "",
                    "",
                    "",
                ]
            )
            continue
        source = entry["source"]
        value = entry["value"]
        sheet.append(
            [
                section_names[field["section"]],
                field["label"],
                field["key"],
                entry["id"],
                value_text(value),
                STATES[entry["state"]],
                entry["certainty"],
                entry["use_scope"],
                source.get("type", "") + ":" + source.get("id", ""),
            ]
        )
        sources.append(
            [
                entry["id"],
                source.get("type", ""),
                source.get("id", ""),
                str(source.get("version", "")),
                source.get("quote", ""),
                entry["use_scope"],
            ]
        )
        if isinstance(value, dict) and field["key"].endswith("events[]"):
            events.append(
                [
                    entry["id"],
                    section_names[field["section"]],
                    value.get("title", ""),
                    value.get("time_raw", ""),
                    value.get("time_precision", ""),
                    value.get("place", ""),
                    value_text(value.get("people", [])),
                    value.get("what", ""),
                    value.get("action", ""),
                    value.get("feelings", ""),
                    value.get("impact") or value.get("result", ""),
                ]
            )
        elif field["key"].endswith("[]"):
            people.append([entry["id"], field["label"], value_text(value)])
    for ws in book:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="244A76")
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 28
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
