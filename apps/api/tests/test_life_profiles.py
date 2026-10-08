from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import func, select

from lifereel_api.core.database import SessionLocal
from lifereel_api.modules.billing.models import Charge
from lifereel_api.modules.interview.profile_models import LifeProfileRevision
from lifereel_api.modules.orchestration.service import execute_turn
from lifereel_api.modules.script.models import ScriptProject


def setup(client):
    person = client.post("/v1/persons", json={"display_name": "资料测试人物"}).json()
    response = client.post("/v1/interviews", json={"subject_id": person["id"]})
    assert response.status_code == 201, response.text
    session = response.json()
    profile = client.get("/v1/life-profiles/subjects/" + person["id"]).json()
    return person, session, profile


def edit(client, profile, changes, request_id=None):
    return client.patch(
        "/v1/life-profiles/" + profile["id"],
        json={
            "expected_version": profile["version_number"],
            "request_id": request_id or str(uuid4()),
            "changes": changes,
        },
    )


def test_person_has_one_life_profile_and_no_chapter(client):
    person, session, profile = setup(client)
    second = client.post("/v1/interviews", json={"subject_id": person["id"]}).json()
    assert second["id"] == session["id"]
    assert second["chapter_id"] is None
    assert second["profile_id"] == profile["id"]
    assert len(profile["fields"]) == 67
    assert len(profile["sections"]) == 12
    assert profile["readiness"]["status"] == "not_ready"
    assert not any(e["field_key"] == "identity.gender" for e in profile["entries"])


def test_edit_is_versioned_idempotent_and_cannot_overwrite_newer(client):
    _, _, profile = setup(client)
    request_id = str(uuid4())
    changes = [{"field_key": "identity.birth_place", "value": "泉州"}]
    first = edit(client, profile, changes, request_id)
    assert first.status_code == 200, first.text
    updated = first.json()
    retry = edit(client, profile, changes, request_id)
    assert retry.status_code == 200
    assert retry.json()["version_number"] == updated["version_number"]
    conflict = edit(client, profile, [{"field_key": "identity.birth_place", "value": "旧值"}])
    assert conflict.status_code == 409
    mismatch = edit(
        client, profile, [{"field_key": "identity.birth_place", "value": "不同请求"}], request_id
    )
    assert mismatch.status_code == 409
    history = client.get(f"/v1/life-profiles/{profile['id']}/history").json()
    assert len(history) == 2


def test_processed_fields_are_not_readiness_and_unknown_is_preserved(client):
    _, _, profile = setup(client)
    response = edit(
        client,
        profile,
        [
            {"field_key": "relationships.applicability", "value": "", "state": "not_applicable"},
            {"field_key": "identity.birth_time", "value": "记不清", "state": "unknown"},
            {"field_key": "work.meaning", "value": "努力"},
        ],
    )
    assert response.status_code == 200, response.text
    assert response.json()["readiness"]["status"] == "not_ready"
    assert response.json()["readiness"]["processed_fields"] == 5


def test_rich_event_is_partial_ready_and_scope_can_be_ready(client):
    _, _, profile = setup(client)
    rich = {
        "title": "工作经历",
        "what": "合成测试经历" * 12,
        "action": "我每天负责记录并认真整理工作中的事情。",
        "impact": "后来我学会了独立处理问题，也开始帮助同事。",
    }
    response = edit(
        client, profile, [{"field_key": "work.events[]", "record_key": "work-1", "value": rich}]
    )
    assert response.status_code == 200, response.text
    updated = response.json()
    assert updated["readiness"]["status"] == "partial_ready"
    scoped = edit(client, updated, [{"field_key": "scope.coverage", "value": {"sections": ["E"]}}])
    assert scoped.json()["readiness"]["status"] == "ready"


def test_private_entries_do_not_leak_to_exports_and_xlsx_is_text(client):
    _, _, profile = setup(client)
    response = edit(
        client,
        profile,
        [
            {
                "field_key": "family.traditions",
                "value": "只留在内部的测试秘密",
                "use_scope": "internal",
            },
            {"field_key": "values.self_description", "value": "=1+1"},
        ],
    )
    assert response.status_code == 200, response.text
    path = f"/v1/life-profiles/{profile['id']}/export"
    md = client.get(path + "?format=md")
    assert "只留在内部" not in md.text
    export = client.get(path + "?format=xlsx")
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content))
    assert len(workbook.sheetnames) == 5
    cells = [c for sheet in workbook for row in sheet for c in row]
    assert not any(c.value == "只留在内部的测试秘密" for c in cells)
    formula = next(c for c in cells if c.value == "=1+1")
    assert formula.data_type == "s"


def test_new_interview_updates_profile_without_script_or_script_charge(client):
    _, session, profile = setup(client)
    response = client.post(
        f"/v1/interviews/{session['id']}/turns",
        json={
            "answer_text": "1988年，我去泉州的工厂做工。",
            "idempotency_key": "life-turn-first",
        },
    )
    assert response.status_code == 202, response.text
    with SessionLocal() as db:
        from lifereel_api.core.config import get_settings

        execute_turn(db, get_settings().default_tenant_id, UUID(response.json()["id"]))
        assert db.scalar(select(func.count()).select_from(ScriptProject)) == 0
        assert (
            db.scalar(select(func.count()).select_from(Charge).where(Charge.kind == "script")) == 0
        )
    updated = client.get(f"/v1/life-profiles/{profile['id']}").json()
    events = [e for e in updated["entries"] if e["field_key"] == "work.events[]"]
    assert len(events) == 1
    assert events[0]["value"]["time_raw"] == "1988年"
    blocked = client.post(
        f"/v1/interviews/{session['id']}/turns",
        json={
            "action": "regenerate_script",
            "idempotency_key": "life-no-script",
        },
    )
    assert blocked.status_code == 409
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(LifeProfileRevision)) == 2


def _rich_profile(client):
    person, session, profile = setup(client)
    value = {
        "title": "工厂经历",
        "time_raw": "1988年",
        "place": "泉州",
        "what": "我在泉州的工厂做工，每天负责登记材料，遇到漏记的单据就重新检查记录。" * 3,
        "action": "我和同事一起核对单据并补齐记录。",
        "impact": "这让我学会耐心，也能独立整理材料。",
    }
    profile = edit(
        client, profile, [{"field_key": "work.events[]", "record_key": "work-1", "value": value}]
    ).json()
    return person, session, profile


def test_dialogue_correction_replaces_event_and_closes_conflict(client):
    from lifereel_api.core.config import get_settings
    from lifereel_api.modules.memory.models import MemoryClaim, MemoryConflict, TimelineAnchor

    person, session, profile = _rich_profile(client)
    tenant = get_settings().default_tenant_id
    with SessionLocal() as db:
        event = next(e for e in profile["entries"] if e["field_key"] == "work.events[]")
        claim = db.scalar(
            select(MemoryClaim).where(MemoryClaim.profile_entry_id == UUID(event["id"]))
        )
        db.add(
            MemoryConflict(
                tenant_id=tenant,
                subject_id=UUID(person["id"]),
                conflict_key="synthetic-profile-timeline",
                claim_ids=[str(claim.id)],
                description="合成旧年份冲突",
                status="open",
            )
        )
        db.commit()
    turn = client.post(
        f"/v1/interviews/{session['id']}/turns",
        json={
            "answer_text": "更正一下，不是1988年，应该是1989年，去漳州的工厂做工。",
            "idempotency_key": "correct-profile-year",
        },
    )
    assert turn.status_code == 202, turn.text
    with SessionLocal() as db:
        execute_turn(db, tenant, UUID(turn.json()["id"]))
        conflicts = list(db.scalars(select(MemoryConflict)))
        assert conflicts and all(c.status == "resolved" for c in conflicts)
        anchors = list(db.scalars(select(TimelineAnchor)))
        assert not any(a.year == 1988 for a in anchors)
    current = client.get(f"/v1/life-profiles/{profile['id']}").json()
    events = [e for e in current["entries"] if e["field_key"] == "work.events[]"]
    assert len(events) == 1
    assert "1989" in str(events[0]["value"]) and "1988" not in str(events[0]["value"])
    assert not any(e["field_key"] == "identity.gender" for e in current["entries"])


def test_saved_book_to_script_and_restricted_source(client):
    person, _, profile = _rich_profile(client)
    created = client.post("/v1/books", json={"subject_id": person["id"]})
    assert created.status_code == 201, created.text
    book = created.json()
    chapter = book["chapters"][0]
    body = "1988年，我在泉州的工厂做工。我和同事一起核对记录，这让我学会耐心。" * 30
    saved = client.patch(
        f"/v1/books/{book['id']}/chapters/{chapter['chapter_id']}",
        json={"expected_version": 0, "title": chapter["title"], "body": body},
    )
    assert saved.status_code == 200, saved.text
    revision = saved.json()["chapters"][0]["current"]
    payload = {
        "subject_id": person["id"],
        "book_revision_ids": [revision["id"]],
        "duration_seconds": 60,
        "idempotency_key": str(uuid4()),
        "adaptation_instructions": "全部镜头不露脸",
    }
    project = client.post("/v1/scripts/generate", json=payload)
    assert project.status_code == 201, project.text
    result = project.json()
    assert result["source_type"] == "book" and len(result["scenes"]) == 2
    assert all(
        shot["visual_constraints"]["face_policy"] == "no_identifiable_faces"
        for shot in result["shots"]
    )
    assert sum(s["duration_seconds"] for s in result["scenes"]) == 60
    assert all(s["chapter_id"] is None for s in result["scenes"])
    replay = client.post("/v1/scripts/generate", json=payload)
    assert replay.json()["id"] == result["id"]
    run = client.post(
        "/v1/production/runs",
        json={
            "project_id": result["id"],
            "scene_id": result["scenes"][0]["id"],
            "audience": "family",
            "provider": "mock",
        },
    )
    assert run.status_code == 201, run.text
    assert run.json()["status"] == "completed"
    event = next(e for e in profile["entries"] if e["field_key"] == "work.events[]")
    restricted = edit(
        client,
        profile,
        [
            {
                "id": event["id"],
                "field_key": event["field_key"],
                "record_key": event["record_key"],
                "value": event["value"],
                "use_scope": "internal",
            }
        ],
    )
    assert restricted.status_code == 200, restricted.text
    assert client.get(f"/v1/books/{book['id']}/export").status_code == 409
    payload["idempotency_key"] = str(uuid4())
    assert client.post("/v1/scripts/generate", json=payload).status_code == 409
    assert (
        client.post(
            "/v1/production/runs",
            json={"project_id": result["id"], "scene_id": result["scenes"][0]["id"]},
        ).status_code
        == 409
    )


def test_memory_failure_retains_committed_profile_and_outbox(client, monkeypatch):
    from lifereel_api.modules.interview import profile_service
    from lifereel_api.modules.jobs.events import OutboxEvent

    _, _, profile = setup(client)

    def fail(*args):
        from lifereel_api.core.errors import ApiError, ErrorCode

        raise ApiError(503, ErrorCode.SERVICE_UNAVAILABLE)

    with monkeypatch.context() as context:
        context.setattr(profile_service, "sync_memory", fail)
        response = edit(client, profile, [{"field_key": "work.meaning", "value": "认真工作"}])
        assert response.status_code == 503
    restored = client.get(f"/v1/life-profiles/{profile['id']}").json()
    assert restored["version_number"] == profile["version_number"] + 1
    with SessionLocal() as db:
        assert (
            db.scalar(
                select(func.count())
                .select_from(OutboxEvent)
                .where(
                    OutboxEvent.event_type == "profile.sync.requested",
                    OutboxEvent.status == "pending",
                )
            )
            == 1
        )


def test_profile_lock_refreshes_stale_identity_map(client):
    import pytest

    from lifereel_api.core.config import get_settings
    from lifereel_api.core.errors import ApiError
    from lifereel_api.modules.interview import profile_service
    from lifereel_api.modules.interview.profile_schemas import ProfilePatch

    _, _, profile = setup(client)
    tenant = get_settings().default_tenant_id
    with SessionLocal() as stale:
        profile_service.get(stale, tenant, UUID(profile["id"]))
        stale.commit()
        assert (
            edit(client, profile, [{"field_key": "work.meaning", "value": "新的手改"}]).status_code
            == 200
        )
        with pytest.raises(ApiError) as error:
            profile_service.patch(
                stale,
                tenant,
                UUID(profile["id"]),
                ProfilePatch(expected_version=profile["version_number"], changes=[]),
            )
        assert error.value.status_code == 409


def test_profile_book_worker_generates_validated_thousand_character_chapter(client, monkeypatch):
    from lifereel_api.core.config import get_settings
    from lifereel_api.modules.book import service as books
    from lifereel_api.modules.book import writing
    from lifereel_api.providers.openai_compatible import OpenAICompatibleClient

    person, _, _ = _rich_profile(client)
    book = client.post("/v1/books", json={"subject_id": person["id"]}).json()
    queued = client.post(f"/v1/books/{book['id']}/generate", json={"idempotency_key": str(uuid4())})
    assert queued.status_code == 202, queued.text
    settings = get_settings()
    monkeypatch.setattr(settings, "llm_provider", "openai-compatible")
    monkeypatch.setattr(settings, "openai_compatible_base_url", "http://qa-model.invalid/v1")
    monkeypatch.setattr(settings, "openai_compatible_api_key", "isolated-qa-model")

    def synthetic(self, system, material):
        import json

        snapshot = json.loads(material)
        claim_id = snapshot["claims"][0]["id"]
        paragraphs = [
            {"text": str(index) + "合成测试记录" * 36, "source_claim_ids": [claim_id]}
            for index in range(5)
        ]
        return {
            "material_sufficient": True,
            "title": "合成书稿",
            "outline": ["记录"],
            "paragraphs": paragraphs,
            "missing_details": [],
        }

    monkeypatch.setattr(OpenAICompatibleClient, "chat_json", synthetic)
    with SessionLocal() as db:
        books.execute(db, settings.default_tenant_id, UUID(queued.json()["job_ids"][0]))
    chapter = client.get(f"/v1/books/{book['id']}").json()["chapters"][0]
    assert writing.MIN_WORDS <= chapter["current"]["word_count"] <= writing.MAX_WORDS
    assert chapter["current"]["author"] == "ai" and not chapter["stale"]


def test_profile_extraction_rejects_an_invented_year(client, monkeypatch):
    import pytest

    from lifereel_api.core.config import get_settings
    from lifereel_api.core.errors import ApiError
    from lifereel_api.modules.interview.profile_extraction import process
    from lifereel_api.providers.openai_compatible import OpenAICompatibleClient

    _, _, profile = setup(client)
    settings = get_settings()
    monkeypatch.setattr(settings, "llm_provider", "openai-compatible")
    monkeypatch.setattr(settings, "openai_compatible_base_url", "http://qa-model.invalid/v1")
    monkeypatch.setattr(settings, "openai_compatible_api_key", "isolated-qa-model")
    monkeypatch.setattr(
        OpenAICompatibleClient,
        "chat_json",
        lambda *args: {
            "changes": [
                {
                    "field_key": "work.events[]",
                    "record_key": "new",
                    "value": {"time_raw": "2001年", "what": "去工厂做工"},
                    "quote": "我去工厂做工。",
                }
            ],
            "next_question": "工作有哪些收获？",
        },
    )
    with SessionLocal() as db, pytest.raises(ApiError) as error:
        process(
            db,
            settings.default_tenant_id,
            UUID(profile["id"]),
            "我去工厂做工。",
            {"type": "test"},
            uuid4(),
        )
    assert error.value.status_code == 502
    assert (
        client.get(f"/v1/life-profiles/{profile['id']}").json()["version_number"]
        == profile["version_number"]
    )


def test_confirmed_legacy_material_correction_removes_old_timeline(client):
    from lifereel_api.core.config import get_settings
    from lifereel_api.modules.memory.models import MemoryClaim, TimelineAnchor

    person = client.post("/v1/persons", json={"display_name": "历史资料回归"}).json()
    tenant = get_settings().default_tenant_id
    with SessionLocal() as db:
        claim = MemoryClaim(
            tenant_id=tenant,
            subject_id=UUID(person["id"]),
            claim_type="recollection",
            claim_text="1988年我去泉州工厂工作。",
            source_quote="1988年我去泉州工厂工作。",
            review_status="verified",
            confidence=1.0,
            extraction_provider="rule",
        )
        db.add(claim)
        db.flush()
        old_id = claim.id
        db.add(
            TimelineAnchor(
                tenant_id=tenant,
                subject_id=UUID(person["id"]),
                claim_id=claim.id,
                year=1988,
                time_text="1988年",
                event_text=claim.claim_text,
                precision="year",
            )
        )
        db.commit()
    profile = client.get("/v1/life-profiles/subjects/" + person["id"]).json()
    legacy = next(e for e in profile["entries"] if e["field_key"] == "legacy.events[]")
    corrected = edit(
        client,
        profile,
        [
            {
                "id": legacy["id"],
                "field_key": legacy["field_key"],
                "record_key": legacy["record_key"],
                "value": {"title": "旧资料更正", "what": "1989年我去漳州工厂工作。"},
            }
        ],
    )
    assert corrected.status_code == 200, corrected.text
    with SessionLocal() as db:
        assert db.get(MemoryClaim, old_id).review_status == "superseded"
        assert not db.scalar(select(TimelineAnchor.id).where(TimelineAnchor.year == 1988))
        assert db.scalar(select(TimelineAnchor.id).where(TimelineAnchor.year == 1989))


def test_generated_book_cannot_reintroduce_old_year_or_invent_gender():
    import pytest

    from lifereel_api.modules.book import writing

    snapshot = {"claims": [{"id": "event", "text": "1989年去工厂做工。",
                            "source_quote": "更正，不是1988年，应该是1989年。"}]}
    draft = {
        "material_sufficient": True,
        "title": "工厂经历",
        "outline": ["经历"],
        "paragraphs": [
            {"text": str(i) + "认真核对记录" * 16, "source_claim_ids": ["event"]}
            for i in range(10)
        ],
        "missing_details": [],
    }
    assert 900 <= writing.validate_output(draft, snapshot)["word_count"] <= 1100
    for invented in ["1988年", "中年女性", "中年男性"]:
        invalid = {**draft, "title": invented + "的工厂经历"}
        with pytest.raises(ValueError, match="unsupported_factual_detail"):
            writing.validate_output(invalid, snapshot)
